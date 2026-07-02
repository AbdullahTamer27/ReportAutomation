"""Damage-section expansion.

A configuration template marks ONE repeatable damage block with two marker
paragraphs containing `{{damage_block_start}}` and `{{damage_block_end}}`.
Everything between them (the three image-placeholder cells, etc.) is the unit
that repeats once per damage point.

For a report with N damage points, the block is deep-copied N times; on copy i
the sentinel `@N` is replaced by i, turning `{{DMG@N_1}}` into `{{DMG1_1}}`,
`{{DMG2_1}}`, … which the image pass then fills. N=0 removes the block entirely
(a clean report). This runs BEFORE the image-placement pass.

Run-splitting note: each affected paragraph's runs are collapsed into one before
substitution, so a `@N` that Word split across runs is still replaced reliably.
"""

import copy
import re

from docx import Document
from docx.oxml.ns import qn

START_TOKEN = "{{damage_block_start}}"
END_TOKEN = "{{damage_block_end}}"
INDEX_SENTINEL = "@N"

# A concrete damage image placeholder, e.g. {{DMG1_2}} (after any expansion).
_DMG_TAG_RE = re.compile(r"\{\{DMG\d+_\d+\}\}")

_W_P = qn("w:p")
_W_R = qn("w:r")
_W_T = qn("w:t")
_WP_DOCPR = qn("wp:docPr")
_PIC_CNVPR = qn("pic:cNvPr")


def _element_text(el):
    return "".join(t.text or "" for t in el.iter(_W_T))


def _has_damage_placeholders(doc):
    """True if the document already holds concrete {{DMGi_j}} image placeholders
    (text or picture Alt Text). These are filled directly by the image pass, so a
    template can carry static damage slots instead of a repeatable marker block."""
    body = doc.element.body
    for p in body.iter(_W_P):
        if _DMG_TAG_RE.search(_element_text(p)):
            return True
    for docPr in body.iter(_WP_DOCPR):
        for attr in ("descr", "name", "title"):
            v = docPr.get(attr)
            if v and _DMG_TAG_RE.search(v):
                return True
    return False


def _subst_index_in_element(el, i):
    """Replace @N -> i in `el`: in paragraph text (runs collapsed so a split
    token is still replaced) AND in any placeholder picture's Alt Text
    (wp:docPr name/descr/title), so {{DMG@N_1}} becomes {{DMG1_1}}, etc.

    Each paragraph is handled by its OWN direct run text (``w:r/w:t``) only —
    never ``.//w:t`` — so a paragraph that anchors a text box does not reach
    down into the box and collapse its runs. Text-box paragraphs live in
    ``w:txbxContent`` and are visited on their own by ``el.iter(_W_P)``, so their
    ``{{ovl_...@N_...}}`` tags are still substituted, but in place."""
    for p in el.iter(_W_P):
        ts = p.findall(_W_R + "/" + _W_T)
        if not ts:
            continue
        joined = "".join(t.text or "" for t in ts)
        if INDEX_SENTINEL not in joined:
            continue
        ts[0].text = joined.replace(INDEX_SENTINEL, str(i))
        for t in ts[1:]:
            t.text = ""

    # Alt Text on placeholder pictures (lives in attributes, not <w:t>).
    for docPr in el.iter(_WP_DOCPR):
        for attr in ("name", "descr", "title"):
            v = docPr.get(attr)
            if v and INDEX_SENTINEL in v:
                docPr.set(attr, v.replace(INDEX_SENTINEL, str(i)))


def _max_drawing_id(body):
    """Highest existing wp:docPr / pic:cNvPr id in the body (0 if none)."""
    mx = 0
    for tag in (_WP_DOCPR, _PIC_CNVPR):
        for el in body.iter(tag):
            try:
                mx = max(mx, int(el.get("id")))
            except (TypeError, ValueError):
                pass
    return mx


def _reassign_drawing_ids(clone, counter):
    """Give every picture in `clone` a fresh unique id so cloned drawings don't
    collide (duplicate ids make Word offer to 'repair' the document)."""
    for tag in (_WP_DOCPR, _PIC_CNVPR):
        for el in clone.iter(tag):
            el.set("id", str(counter[0]))
            counter[0] += 1


def expand_damage_blocks(doc, damage_count):
    """Expand the marked damage block in `doc` `damage_count` times in place.

    Returns True if the markers were found (block expanded/removed), else False.
    """
    body = doc.element.body

    start = end = None
    for el in list(body):
        if el.tag == _W_P:
            txt = _element_text(el)
            if start is None and START_TOKEN in txt:
                start = el
            elif start is not None and END_TOKEN in txt:
                end = el
                break

    if start is None or end is None:
        return False

    # Elements strictly between the markers, in document order.
    block, collecting = [], False
    for el in list(body):
        if el is start:
            collecting = True
            continue
        if el is end:
            break
        if collecting:
            block.append(el)

    # Insert N clones just before the end marker (preserves order). Cloned
    # pictures get fresh unique ids so Word doesn't flag duplicate drawing ids.
    id_counter = [_max_drawing_id(body) + 1]
    for i in range(1, int(damage_count) + 1):
        for el in block:
            clone = copy.deepcopy(el)
            _subst_index_in_element(clone, i)
            _reassign_drawing_ids(clone, id_counter)
            end.addprevious(clone)

    # Remove the original block + both markers.
    for el in block:
        body.remove(el)
    body.remove(start)
    body.remove(end)
    return True


def expand_in_file(path, damage_count, progress=None, review=None):
    """Open `path`, expand the damage block `damage_count` times, save in place."""
    log = progress or (lambda m: None)
    rev = review or (lambda m: None)

    doc = Document(path)
    found = expand_damage_blocks(doc, damage_count)
    doc.save(path)

    if found:
        log(f"Damage sections: block expanded x{int(damage_count)}.")
    elif damage_count and int(damage_count) > 0:
        # No repeatable block — but if the template carries static {{DMGi_j}}
        # placeholders, the image pass fills them, so don't cry wolf.
        if _has_damage_placeholders(doc):
            log(f"Damage sections: no {START_TOKEN}/{END_TOKEN} block, but static "
                f"{{{{DMGi_j}}}} placeholders are present — the image pass fills them.")
        else:
            rev(f"⚠ Number of damages = {int(damage_count)}, but the template has no "
                f"{START_TOKEN}/{END_TOKEN} markers and no {{{{DMGi_j}}}} placeholders — "
                f"no damage pictures were added.")
    else:
        log("Damage sections: no block markers (none requested).")
    return found
